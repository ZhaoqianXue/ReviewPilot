#!/usr/bin/env python3
"""Merge verify-pass results into the redo xlsx for the 25 draft-only rows.
Only overwrites a row if a valid verified final exists; logs U changes vs draft.
Usage: python3 merge_verify.py verify_results.json
"""
import sys, json
import openpyxl
from openpyxl.utils import column_index_from_string

REDO = '/Users/zhaoqianxue/Desktop/UA/ReviewPilot/agent_skill/agent_skills_infor_claude_redo.xlsx'
COLS = ['L','M','N','O','P','Q','R','S','T','U']
CIDX = {c: column_index_from_string(c) for c in COLS}

def main(path):
    vres = json.load(open(path))
    drafts = json.load(open('/Users/zhaoqianxue/Desktop/UA/ReviewPilot/tmp/redo_extract/workflow_results.json'))
    wb = openpyxl.load_workbook(REDO)
    ws = wb.active

    updated, skipped, uchanges = [], [], []
    for rid, rec in vres.items():
        n = int(rid.replace('row',''))
        pr = n + 1
        if ws.cell(pr,1).value != rid:
            skipped.append(f'{rid}: A-col mismatch'); continue
        fin = rec.get('final')
        if not fin or not fin.get('U'):
            skipped.append(f'{rid}: no verified final (kept draft)'); continue
        old_u = drafts[rid]['draft']['U']
        new_u = fin['U']
        if new_u != old_u:
            uchanges.append(f'{rid}: U {old_u} -> {new_u}  | {fin.get("notes","")[:160]}')
        for c in COLS:
            v = fin.get(c)
            if v is None or str(v).strip()=='':
                v = 'N/A'
            ws.cell(pr, CIDX[c]).value = str(v).strip()
        updated.append(rid)

    wb.save(REDO)

    # validate
    wb2 = openpyxl.load_workbook(REDO); ws2 = wb2.active
    blanks=[]; aerr=[]
    from collections import Counter
    dist=Counter()
    for n in range(51,101):
        pr=n+1
        if ws2.cell(pr,1).value != f'row{n}': aerr.append(pr)
        for c in COLS:
            v=ws2.cell(pr,CIDX[c]).value
            if v is None or str(v).strip()=='': blanks.append(f'row{n}:{c}')
        dist[ws2.cell(pr,CIDX['U']).value]+=1

    print('Updated rows:', len(updated))
    print('Skipped (kept draft):', skipped if skipped else 'NONE')
    print('U changes by verify:', uchanges if uchanges else 'NONE')
    print('Blanks:', blanks if blanks else 'NONE')
    print('A-col errors:', aerr if aerr else 'NONE')
    print('Final U distribution (all 50):', dict(dist))

if __name__=='__main__':
    main(sys.argv[1])
