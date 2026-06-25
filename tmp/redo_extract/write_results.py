#!/usr/bin/env python3
"""Write workflow L-U results into the redo xlsx for row51-row100, with validation.
Usage: python3 write_results.py results.json
results.json = the workflow return object {rid: {draft, final, flag}, ...}
"""
import sys, json
import openpyxl
from openpyxl.utils import column_index_from_string

REDO = '/Users/zhaoqianxue/Desktop/UA/ReviewPilot/agent_skill/agent_skills_infor_claude_redo.xlsx'
COLS = ['L','M','N','O','P','Q','R','S','T','U']
CIDX = {c: column_index_from_string(c) for c in COLS}

def pick(rec):
    """Prefer verified final; fall back to draft."""
    final = rec.get('final')
    draft = rec.get('draft')
    src = final if (final and final.get('U')) else draft
    return src

def main(path):
    data = json.load(open(path))
    wb = openpyxl.load_workbook(REDO)
    ws = wb.active

    problems = []
    summary = []
    for n in range(51, 101):
        rid = f'row{n}'
        pr = n + 1  # physical row
        # sanity: A column must still be rowNN
        a = ws.cell(pr, 1).value
        if a != rid:
            problems.append(f'{rid}: A-col mismatch at phys {pr} -> {a!r}')
            continue
        rec = data.get(rid)
        if not rec:
            problems.append(f'{rid}: NO RESULT in json')
            continue
        src = pick(rec)
        if not src:
            problems.append(f'{rid}: no draft/final payload')
            continue
        for c in COLS:
            v = src.get(c)
            if v is None or str(v).strip() == '':
                problems.append(f'{rid}: blank {c}')
                v = 'N/A'
            ws.cell(pr, CIDX[c]).value = str(v).strip()
        used = 'final' if (rec.get('final') and rec['final'].get('U')) else 'draft'
        flipped = rec.get('final',{}).get('changed_u') if rec.get('final') else None
        summary.append((rid, src.get('U'), used, flipped, src.get('confidence')))

    wb.save(REDO)

    # ---- validation pass on the saved file ----
    wb2 = openpyxl.load_workbook(REDO)
    ws2 = wb2.active
    blanks = []
    aerr = []
    for n in range(51,101):
        pr = n+1
        if ws2.cell(pr,1).value != f'row{n}':
            aerr.append((pr, ws2.cell(pr,1).value))
        for c in COLS:
            v = ws2.cell(pr, CIDX[c]).value
            if v is None or str(v).strip()=='':
                blanks.append(f'row{n}:{c}')

    print('=== WRITE SUMMARY ===')
    for rid,u,used,flip,conf in summary:
        fl = ' U-FLIPPED' if flip else ''
        print(f'{rid:7} U={u:14} [{used},{conf}]{fl}')
    print()
    print('Rows written:', len(summary))
    print('Blanks after save:', blanks if blanks else 'NONE')
    print('A-col errors:', aerr if aerr else 'NONE')
    print('Problems during write:', problems if problems else 'NONE')
    # U distribution
    from collections import Counter
    dist = Counter(u for _,u,_,_,_ in summary)
    print('U distribution:', dict(dist))

if __name__ == '__main__':
    main(sys.argv[1])
